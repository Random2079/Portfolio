"""Fund TER / fees from Bank of Russia PIF showcase Excel (free).

Primary metric for the review slot: УК fee + max total expenses (closest to TER).
Universe cached in SQLite — first hit downloads xlsx, then lookups are local.
"""

from __future__ import annotations

import logging
import re
import time
from dataclasses import asdict, dataclass
from io import BytesIO
from typing import Any, Optional

import requests
from sqlalchemy.orm import Session

from portfolio_news.db import FundTerUniverseCache

log = logging.getLogger(__name__)

CBR_SHOWCASE_URL = "https://cbr.ru/RSCI/data_showcase/"
CBR_XLSX_FALLBACK = (
    "https://cbr.ru/Content/Document/File/193443/%20mutual_fund_data.xlsx"
)
_HEADERS = {
    "User-Agent": "PortfolioNews/0.2 (+local; personal monitor)",
    "Accept": "*/*",
}
# CBR updates ~monthly; keep a week so morning terminal is always local.
_CACHE_FRESH_SEC = 7 * 24 * 3600
_PCT_RE = re.compile(
    r"(?P<num>\d+(?:[.,]\d+)?)\s*%",
    re.I,
)


@dataclass
class FundTerFacts:
    isin: str
    name: str = ""
    uk_fee_pct: Optional[float] = None
    max_expense_pct: Optional[float] = None
    uk_fee_raw: str = ""
    max_expense_raw: str = ""
    sheet_as_of: str = ""  # «по состоянию на …» from workbook

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict[str, Any]) -> "FundTerFacts":
        return cls(
            isin=str(d.get("isin") or "").strip().upper(),
            name=str(d.get("name") or ""),
            uk_fee_pct=_opt_float(d.get("uk_fee_pct")),
            max_expense_pct=_opt_float(d.get("max_expense_pct")),
            uk_fee_raw=str(d.get("uk_fee_raw") or ""),
            max_expense_raw=str(d.get("max_expense_raw") or ""),
            sheet_as_of=str(d.get("sheet_as_of") or ""),
        )

    def display(self) -> Optional[str]:
        """Short slot text — not advice."""
        bits: list[str] = []
        if self.uk_fee_pct is not None:
            bits.append(f"УК {_fmt_pct(self.uk_fee_pct)}")
        if self.max_expense_pct is not None:
            bits.append(f"макс. {_fmt_pct(self.max_expense_pct)}")
        if bits:
            return " · ".join(bits)
        # raw fallback if % parse failed but text exists
        raw = (self.uk_fee_raw or self.max_expense_raw or "").strip()
        return raw[:80] if raw and raw.lower() not in {"не предусмотрено", "-"} else None


def _opt_float(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _fmt_pct(n: float) -> str:
    s = f"{n:.2f}".rstrip("0").rstrip(".")
    return s.replace(".", ",") + "%"


def parse_fee_pct(raw: Any) -> Optional[float]:
    """Extract first percent from CBR cell text (e.g. «0,7% от среднегодовой СЧА»)."""
    s = str(raw or "").strip()
    if not s:
        return None
    low = s.lower()
    if low in {"не предусмотрено", "-", "—", "–", "n/a"}:
        return None
    m = _PCT_RE.search(s.replace(" ", ""))
    if not m:
        m = _PCT_RE.search(s)
    if not m:
        return None
    try:
        return float(m.group("num").replace(",", "."))
    except ValueError:
        return None


def discover_cbr_xlsx_url(html: str = "") -> str:
    """Pick .xlsx href from showcase page; fallback to known File id."""
    if not html:
        try:
            resp = requests.get(CBR_SHOWCASE_URL, headers=_HEADERS, timeout=30)
            resp.raise_for_status()
            html = resp.text
        except Exception as exc:  # noqa: BLE001
            log.warning("cbr showcase page failed: %s", exc)
            return CBR_XLSX_FALLBACK
    m = re.search(
        r'href="([^"]*Document/File/[^"]+\.xlsx[^"]*)"',
        html,
        re.I,
    )
    if not m:
        return CBR_XLSX_FALLBACK
    href = m.group(1).replace("&amp;", "&")
    if href.startswith("http"):
        return href
    if href.startswith("/"):
        return "https://cbr.ru" + href
    return CBR_XLSX_FALLBACK


def parse_cbr_pif_xlsx(data: bytes) -> tuple[dict[str, FundTerFacts], str]:
    """Pure-ish xlsx parser. Returns (by_isin, sheet_as_of)."""
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    sheet_as_of = ""
    header_idx: Optional[int] = None
    col: dict[str, int] = {}
    out: dict[str, FundTerFacts] = {}

    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        vals = list(row)
        if ri <= 2:
            blob = " ".join(str(v) for v in vals if v is not None)
            if "по состоянию" in blob.lower():
                sheet_as_of = blob.strip()
            continue
        if header_idx is None:
            joined = " ".join(str(v).lower() for v in vals if v is not None)
            if "isin" in joined and "вознаграждение" in joined:
                header_idx = ri
                for i, v in enumerate(vals):
                    h = str(v or "").lower()
                    if "isin" in h:
                        col["isin"] = i
                    elif "наименование паевого" in h or (
                        "наименование" in h and "фонд" in h and "управл" not in h
                    ):
                        col["name"] = i
                    elif "вознаграждение управляющей компании" in h and "успех" not in h and "результат" not in h:
                        col["uk"] = i
                    elif "максимальный совокупный размер расходов" in h:
                        col["max"] = i
                continue
            continue
        # skip numeric sub-header row (1,2,3,...)
        if ri == header_idx + 1:
            first = vals[0] if vals else None
            if isinstance(first, (int, float)) or str(first).strip().isdigit():
                continue
        isin_i = col.get("isin")
        if isin_i is None or isin_i >= len(vals):
            continue
        isin = str(vals[isin_i] or "").strip().upper()
        if not isin.startswith("RU") or len(isin) < 10:
            continue
        name_i = col.get("name")
        uk_i = col.get("uk")
        max_i = col.get("max")
        name = str(vals[name_i] or "").strip() if name_i is not None and name_i < len(vals) else ""
        uk_raw = str(vals[uk_i] or "").strip() if uk_i is not None and uk_i < len(vals) else ""
        max_raw = str(vals[max_i] or "").strip() if max_i is not None and max_i < len(vals) else ""
        out[isin] = FundTerFacts(
            isin=isin,
            name=name,
            uk_fee_pct=parse_fee_pct(uk_raw),
            max_expense_pct=parse_fee_pct(max_raw),
            uk_fee_raw=uk_raw,
            max_expense_raw=max_raw,
            sheet_as_of=sheet_as_of,
        )
    wb.close()
    return out, sheet_as_of


def fetch_cbr_fund_ter_universe() -> dict[str, FundTerFacts]:
    url = discover_cbr_xlsx_url()
    try:
        resp = requests.get(url, headers=_HEADERS, timeout=60)
        resp.raise_for_status()
        by_isin, as_of = parse_cbr_pif_xlsx(resp.content)
        log.info("cbr pif ter: %s funds (as_of=%s)", len(by_isin), as_of or "?")
        return by_isin
    except Exception as exc:  # noqa: BLE001
        log.warning("cbr pif xlsx failed (%s): %s", url, exc)
        return {}


def _now() -> float:
    return time.time()


def load_universe_cache(session: Session) -> Optional[tuple[dict[str, FundTerFacts], float, str]]:
    import json

    row = session.get(FundTerUniverseCache, 1)
    if row is None or not (row.payload_json or "").strip():
        return None
    try:
        data = json.loads(row.payload_json)
    except json.JSONDecodeError:
        return None
    if not isinstance(data, dict):
        return None
    out = {
        k: FundTerFacts.from_dict(v)
        for k, v in data.items()
        if isinstance(v, dict) and k
    }
    return out, float(row.updated_at or 0.0), str(row.source_as_of or "")


def save_universe_cache(
    session: Session, by_isin: dict[str, FundTerFacts], *, source_as_of: str = ""
) -> None:
    import json

    payload = {isin: f.to_dict() for isin, f in by_isin.items()}
    row = session.get(FundTerUniverseCache, 1)
    if row is None:
        row = FundTerUniverseCache(id=1)
        session.add(row)
    row.payload_json = json.dumps(payload, ensure_ascii=False)
    row.updated_at = _now()
    row.ok = 1 if payload else 0
    row.source_as_of = source_as_of or ""
    session.commit()


def universe_fresh(updated_at: float, *, max_age: float = _CACHE_FRESH_SEC) -> bool:
    return updated_at > 0 and (_now() - updated_at) < max_age


def get_fund_ter_map(
    session: Session, *, force: bool = False
) -> tuple[dict[str, FundTerFacts], str]:
    """Cache-first full ISIN map. Returns (map, as_of_hint)."""
    from datetime import datetime, timedelta, timezone

    as_of = ""
    if not force:
        cached = load_universe_cache(session)
        if cached and universe_fresh(cached[1]) and len(cached[0]) >= 50:
            raw, upd, _src = cached
            as_of = datetime.fromtimestamp(upd, tz=timezone(timedelta(hours=5))).strftime(
                "%Y-%m-%d %H:%M"
            )
            return raw, as_of

    fetched = fetch_cbr_fund_ter_universe()
    if fetched:
        sheet_as_of = next(iter(fetched.values())).sheet_as_of if fetched else ""
        try:
            save_universe_cache(session, fetched, source_as_of=sheet_as_of)
        except Exception as exc:  # noqa: BLE001
            log.warning("fund ter cache save failed: %s", exc)
        as_of = datetime.now(timezone(timedelta(hours=5))).strftime("%Y-%m-%d %H:%M")
        return fetched, as_of

    cached = load_universe_cache(session)
    if cached:
        raw, upd, _src = cached
        as_of = datetime.fromtimestamp(upd, tz=timezone(timedelta(hours=5))).strftime(
            "%Y-%m-%d %H:%M"
        )
        return raw, as_of
    return {}, ""


def resolve_fund_ter(
    by_isin: dict[str, FundTerFacts], isin: str
) -> Optional[FundTerFacts]:
    tid = (isin or "").strip().upper()
    if not tid:
        return None
    return by_isin.get(tid)
